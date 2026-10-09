"""Regenera el adjunto editable desde docs/guia-workflow-ia.md.

Requiere python-docx y Pillow. Ejecutar desde la raíz del repositorio.
Las figuras son ejemplos; las respuestas y tablas siguen siendo editables.
"""
from pathlib import Path
import re
import tempfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'api/app/recursos/guia-workflow-ia.docx'


def diagram(path, micro=False):
    im = Image.new('RGB', (1500, 900 if micro else 240), 'white')
    d = ImageDraw.Draw(im)
    candidates = ['/System/Library/Fonts/Supplemental/Arial.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']
    font = next((ImageFont.truetype(p, 25) for p in candidates if Path(p).exists()), ImageFont.load_default(size=25))
    def label(x, y, t):
        d.multiline_text((x, y), t, fill='#152536', font=font, anchor='mm', align='center', spacing=5)
    def box(x, y, w, h, t, diamond=False):
        if diamond:
            d.polygon([(x,y-h/2),(x+w/2,y),(x,y+h/2),(x-w/2,y)], fill='#eef2f5', outline='#536777', width=3)
        else:
            d.rounded_rectangle((x-w/2,y-h/2,x+w/2,y+h/2), radius=14, fill='#eef2f5', outline='#536777', width=3)
        label(x,y,t)
    def arrow(points, text=None, at=None):
        d.line(points, fill='#536777', width=3)
        x,y=points[-1]; px,py=points[-2]
        if x==px:
            sign=1 if y>py else -1
            d.polygon([(x,y),(x-8,y-14*sign),(x+8,y-14*sign)],fill='#536777')
        else:
            sign=1 if x>px else -1
            d.polygon([(x,y),(x-14*sign,y-8),(x-14*sign,y+8)],fill='#536777')
        if text: label(*at,text)
    if not micro:
        texts=['Recibir\npedido','Validar\npedido','Registrar\npedido','Preparar\ny entregar','Facturar\ny cobrar']
        for i,t in enumerate(texts):
            x=145+300*i
            box(x,120,250,120,t)
            if i<4: arrow([(x+125,120),(x+175,120)])
    else:
        box(440,60,360,80,'Pedido recibido')
        box(440,190,540,100,'Revisar producto,\ncantidad y precio')
        box(440,375,510,200,'¿Datos correctos\ny completos?',True)
        box(440,575,360,80,'Pedido validado')
        box(1110,375,420,90,'Solicitar corrección')
        box(1110,575,500,190,'¿Se recibe\nla corrección?',True)
        box(1110,815,620,90,'Informar pedido pendiente\ny escalar')
        arrow([(440,100),(440,140)])
        arrow([(440,240),(440,275)])
        arrow([(440,475),(440,535)],'Sí',(480,505))
        arrow([(695,375),(900,375)],'No',(790,345))
        arrow([(1110,420),(1110,480)])
        arrow([(1110,670),(1110,770)],'No, vence el plazo',(1250,720))
        arrow([(1360,575),(1450,575),(1450,190),(710,190)],'Sí',(1410,255))
    im.save(path)


def text(p, value):
    for i, part in enumerate(re.split(r'\*\*(.*?)\*\*', value)):
        r=p.add_run(part); r.bold=bool(i%2)


def build():
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.65)
    sec.left_margin=sec.right_margin=Inches(.65)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','List Bullet']:
        st=doc.styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
    for st in doc.styles:
        for border in st.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.space_after=Pt(7)
    doc.styles['Normal'].paragraph_format.line_spacing=1.06
    doc.styles['Title'].font.size=Pt(24)
    doc.styles['Heading 1'].font.size=Pt(16)
    doc.styles['Heading 2'].font.size=Pt(13)
    # Forms are unprotected: respondents can type answers and add table rows.
    doc.core_properties.title='Guía para proponer un workflow de IA'
    doc.core_properties.author='IA Nexo'
    lines=(ROOT/'docs/guia-workflow-ia.md').read_text().splitlines()
    i=0; diagrams=0
    with tempfile.TemporaryDirectory() as tmp:
        while i<len(lines):
            line=lines[i].strip()
            if not line: i+=1; continue
            if line.startswith('```'):
                j=i+1
                while j<len(lines) and not lines[j].startswith('```'): j+=1
                path=Path(tmp)/f'diagram-{diagrams}.png'; diagram(path,diagrams==1)
                p=doc.add_paragraph(); p.add_run().add_picture(str(path),width=Inches(7.1))
                for shape in p._p.xpath('.//wp:docPr'):
                    shape.set('descr','Ejemplo de flujograma micro de validación de pedidos' if diagrams else 'Ejemplo de flujograma macro de pedidos')
                diagrams+=1; i=j+1; continue
            if line.startswith('|'):
                rows=[]
                while i<len(lines) and lines[i].strip().startswith('|'):
                    cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                    if not all(re.fullmatch(r'[-: ]+',c) for c in cells): rows.append(cells)
                    i+=1
                table=doc.add_table(rows=0,cols=len(rows[0])); table.style='Table Grid'
                for n,row in enumerate(rows):
                    cells=table.add_row().cells
                    for cell,value in zip(cells,row):
                        p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(5); p.paragraph_format.space_before=Pt(5)
                        text(p,value.replace('min-persona/caso', 'min-persona por caso'))
                        for r in p.runs: r.font.size=Pt(9)
                        if n==0:
                            for r in p.runs: r.bold=True
                            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E8EDF2'); cell._tc.get_or_add_tcPr().append(shade)
                    trPr=table.rows[-1]._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
                    if n==0: trPr.append(OxmlElement('w:tblHeader'))
                doc.add_paragraph().paragraph_format.space_after=Pt(0)
                continue
            if line.startswith('#'):
                level=len(line)-len(line.lstrip('#')); title=line.lstrip('# ').replace('¿','').replace('?','').replace(':','')
                heading=doc.add_paragraph(title, 'Title' if level==1 else ('Heading 1' if level==2 else 'Heading 2'))
                if title.startswith(('4.', 'Fuentes de información','Mapa de flujo','5.')): heading.paragraph_format.page_break_before=True
            elif line.startswith('- '): text(doc.add_paragraph(style='List Bullet'),line[2:])
            else: text(doc.add_paragraph(),line)
            # Simple answer space in the narrative sections, without adding people tables.
            if line.startswith(('¿De qué **','¿Qué proceso realiza','Dentro de ese proceso,','**Enlace a la carpeta:**')):
                doc.add_paragraph('Respuesta: __________________________________________________')
            i+=1
    OUT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUT)
    print(OUT)

if __name__=='__main__': build()
